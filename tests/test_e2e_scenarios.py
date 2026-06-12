from __future__ import annotations

import io
import json
import time
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.demo import DemoParserCore
from app.jobs import JobManager
from app.main import app
from app.storage import SQLiteStore


def configure(tmp_path: Path, delay: float = 0.001) -> TestClient:
    import app.main as main

    store = SQLiteStore(tmp_path / "e2e.db")
    store.initialize()
    main.store = store
    main.manager = JobManager(store, core_factory=lambda: DemoParserCore(delay))
    return TestClient(app)


def wait_for(client: TestClient, job_id: str, statuses: set[str]) -> dict:
    for _ in range(300):
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in statuses:
            return job
        time.sleep(0.01)
    raise AssertionError(f"job {job_id} did not reach {statuses}")


def create(client: TestClient, query: str, area: dict, target: int) -> str:
    response = client.post(
        "/api/jobs",
        json={"query": query, "area": area, "params": {"target_businesses_count": target}},
    )
    assert response.status_code == 202
    return response.json()["id"]


def test_e2e_small_bbox_narrow_query_and_all_exports(tmp_path: Path) -> None:
    client = configure(tmp_path)
    job_id = create(
        client,
        "кофейни",
        {"type": "bbox", "coordinates": [37.61, 55.75, 37.62, 55.76]},
        5,
    )
    job = wait_for(client, job_id, {"completed"})
    page = client.get(f"/api/jobs/{job_id}/results?limit=100").json()
    assert job["found"] == page["total"] == 5
    assert len({item["dedupe_key"] for item in page["items"]}) == 5
    assert all(item["latitude"] and item["longitude"] for item in page["items"])

    json_export = client.get(f"/api/jobs/{job_id}/export?format=json").content
    csv_export = client.get(f"/api/jobs/{job_id}/export?format=csv").content
    xlsx_export = client.get(f"/api/jobs/{job_id}/export?format=xlsx").content
    assert len(json.loads(json_export)) == 5
    assert b"name" in csv_export
    assert load_workbook(io.BytesIO(xlsx_export)).active.max_row == 6


def test_e2e_medium_bbox_popular_category_and_repeat_run(tmp_path: Path) -> None:
    client = configure(tmp_path)
    area = {"type": "bbox", "coordinates": [37.50, 55.65, 37.80, 55.85]}
    first = create(client, "аптеки", area, 20)
    second = create(client, "аптеки", area, 20)
    assert wait_for(client, first, {"completed"})["found"] == 20
    assert wait_for(client, second, {"completed"})["found"] == 20
    for job_id in (first, second):
        page = client.get(f"/api/jobs/{job_id}/results?limit=100").json()
        assert page["total"] == 20
        assert len({item["yandex_id"] for item in page["items"]}) == 20


def test_e2e_complex_polygon_stop_and_partial_export(tmp_path: Path) -> None:
    client = configure(tmp_path, delay=0.03)
    polygon = {
        "type": "polygon",
        "coordinates": [
            [37.50, 55.65],
            [37.80, 55.65],
            [37.72, 55.76],
            [37.80, 55.85],
            [37.50, 55.85],
            [37.58, 55.76],
            [37.50, 55.65],
        ],
    }
    job_id = create(client, "рестораны", polygon, 40)
    for _ in range(100):
        page = client.get(f"/api/jobs/{job_id}/results?limit=100").json()
        if page["total"] >= 2:
            break
        time.sleep(0.01)
    client.post(f"/api/jobs/{job_id}/stop")
    job = wait_for(client, job_id, {"stopped"})
    page = client.get(f"/api/jobs/{job_id}/results?limit=100").json()
    assert 1 <= page["total"] < 40
    assert job["found"] == page["total"]
    assert client.get(f"/api/jobs/{job_id}/export?format=json").content != b"[]"

